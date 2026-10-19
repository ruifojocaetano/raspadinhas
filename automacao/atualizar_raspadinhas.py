# -*- coding: utf-8 -*-
"""
Atualiza automaticamente a lista de Raspadinhas dos Jogos Santa Casa.

Origem:
https://www.jogossantacasa.pt/web/JogarLotInst

O programa:
1. Abre a página das Raspadinhas.
2. Descobre automaticamente todos os links "Saiba Mais" (e respetivo contentId).
3. Abre cada página de detalhe.
4. Extrai jogo, nome, preço, prémio máximo, emissão e data de lançamento.
5. Cria/atualiza Raspadinhas.xlsx.

Não é necessário indicar manualmente nenhum contentId.
"""

import re
import time
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


URL_LISTA = "https://www.jogossantacasa.pt/web/JogarLotInst"
FICHEIRO_EXCEL = "Raspadinhas.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/141.0 Safari/537.36"
    )
}


def limpar_texto(texto):
    return re.sub(r"\s+", " ", texto or "").strip()


def pedir_pagina(session, url):
    resposta = session.get(url, timeout=30)
    resposta.raise_for_status()
    # O site declara "ISO-8859-1" no cabeçalho, mas na prática envia os
    # bytes em Windows-1252 (comum em sites PT antigos). Isto é o que
    # corrompe o símbolo "€" e os acentos se deixarmos o requests adivinhar.
    resposta.encoding = "cp1252"
    return BeautifulSoup(resposta.text, "html.parser")


def limpar_nome(nome):
    """Remove restos de '(data de lançamento ...)' e 'Saiba Mais' que por
    vezes ficam colados ao nome por causa da estrutura da página de lista."""
    nome = re.sub(r"\(data de lançamento.*?\)", "", nome, flags=re.I)
    nome = re.sub(r"Saiba Mais", "", nome, flags=re.I)
    return limpar_texto(nome)


def extrair_jogos_da_lista(soup):
    """
    Só recolhe o contentId/URL de cada jogo a partir dos links de detalhe.
    Jogo, Nome e Data são extraídos de forma fiável mais tarde, a partir
    da própria página de detalhe (evita apanhar texto de jogos vizinhos).
    """
    jogos = []
    vistos = set()

    links = soup.select('a[href*="caracteristicasOutras"]')

    for link in links:
        href = link.get("href", "").strip()
        if not href:
            continue

        url = urljoin(URL_LISTA, href)

        match_id = re.search(r"contentId=([0-9.]+)", url, re.I)
        if not match_id:
            continue

        content_id = match_id.group(1)

        if content_id in vistos:
            continue
        vistos.add(content_id)

        jogos.append({
            "Jogo": None,
            "Nome": "",
            "Data lançamento": "",
            "contentId": content_id,
            "URL": url,
        })

    return jogos


def texto_para_numero(texto):
    """Converte um valor extraído (ex.: '2', '2,50', '1.500') num float,
    tratando '.' como separador de milhares e ',' como separador decimal
    (formato português). Devolve None se não for possível converter."""
    if not texto:
        return None
    limpo = texto.replace(".", "").replace(",", ".")
    try:
        return float(limpo)
    except ValueError:
        return None


def extrair_valor(texto, padroes):
    """Tenta vários padrões (o site usa mais do que um template) e devolve
    o primeiro valor encontrado."""
    if isinstance(padroes, str):
        padroes = [padroes]
    for padrao in padroes:
        match = re.search(padrao, texto, re.I)
        if match:
            return limpar_texto(match.group(1))
    return ""


def extrair_detalhes(session, jogo):
    soup = pedir_pagina(session, jogo["URL"])
    texto = limpar_texto(soup.get_text(" ", strip=True))

    # --- Título / Nome -----------------------------------------------
    # NOTA: soup.find("h1") apanha o logótipo "Jogos Santa Casa" no topo
    # da página (é o h1 real do site), NÃO o título do jogo — esse está
    # num <h4>. Por isso procuramos o padrão "Jogo N - Nome" diretamente
    # no texto da página (aparece no breadcrumb e no cabeçalho da secção).
    jogo["Jogo"] = None
    jogo["Nome"] = ""

    match_titulo = re.search(
        r"Jogo\s+(\d+)\s*[-–—]\s*(.+?)\s*\(data de lançamento",
        texto,
        re.I,
    )
    if not match_titulo:
        # Sem a parte "(data de lançamento...)" -> apanha até à próxima secção conhecida
        match_titulo = re.search(
            r"Jogo\s+(\d+)\s*[-–—]\s*(.+?)(?=\s+Emiss[aã]o[^\d]|\s+Raspadinha\s|\s*$)",
            texto,
            re.I,
        )

    if match_titulo:
        jogo["Jogo"] = int(match_titulo.group(1))
        jogo["Nome"] = limpar_nome(match_titulo.group(2))
    else:
        # Alternativa: "Jogo: 266 · Nome do Jogo: Mini Pé-De-Meia"
        match_num = re.search(r"Jogo\s*:\s*(\d+)", texto, re.I)
        match_nome = re.search(r"Nome do Jogo\s*:\s*(.*?)\s*(?:Preço|Emiss[aã]o|Pr[eé]mio)", texto, re.I)
        if match_num:
            jogo["Jogo"] = int(match_num.group(1))
        if match_nome:
            jogo["Nome"] = limpar_nome(match_nome.group(1))

    # --- Emissão --------------------------------------------------------
    # Em vez de tentar prever todos os separadores/traços possíveis,
    # ignoramos tudo o que não for dígito entre "Emissão" e o número —
    # assim não interessa se o site usa "-", "–", ":" ou outro texto.
    jogo["Emissão"] = extrair_valor(texto, [
        r"Emiss[aã]o[^\d]*([\d\.\s]+?)\s*(?:de\s+)?bilhetes",
    ])

    # --- Preço (o símbolo € é um ícone, não texto — por isso não se procura) ---
    # Mesma lógica: aceita "Preço do bilhete", "Preço do bilhete para o
    # apostador", "Preço por Bilhete", etc. — ignora tudo até ao € ou ao número.
    preco_texto = extrair_valor(texto, [
        r"Preço (?:do|por) [Bb]ilhete[^€\d]*€?\s*([\d.,]+)",
    ])
    jogo["Preço"] = texto_para_numero(preco_texto)

    # --- Prémio máximo ----------------------------------------------------
    jogo["Prémio máximo"] = extrair_valor(texto, [
        r"Pr[eé]mio m[aá]ximo[^€\d]*€?\s*([\d.,]+)",
    ])

    # --- Data de lançamento -----------------------------------------------
    # Captura diretamente o padrão de data ("10 de maio de 2026"), em vez
    # de tentar adivinhar onde a frase termina.
    data_detalhe = extrair_valor(texto, [
        r"Data de lançamento[^\d]*(\d{1,2}\s+de\s+\w+\s+de\s+\d{4})",
    ])
    if data_detalhe:
        jogo["Data lançamento"] = data_detalhe

    return jogo


def criar_excel(jogos, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Raspadinhas"

    colunas = [
        "Jogo",
        "Nome",
        "Preço",
        "Prémio máximo",
        "Emissão",
        "Data lançamento",
        "contentId",
        "URL",
        "Atualizado em",
    ]

    ws.append(colunas)

    agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    def chave_ordenacao(j):
        # 1º critério: Preço ascendente (jogos sem preço vão para o fim)
        # 2º critério: Jogo ascendente (jogos sem número vão para o fim)
        preco = j.get("Preço")
        preco_chave = preco if preco is not None else float("inf")
        numero_jogo = j.get("Jogo")
        jogo_chave = numero_jogo if numero_jogo is not None else float("inf")
        return (preco_chave, jogo_chave)

    for jogo in sorted(jogos, key=chave_ordenacao):
        linha = [
            jogo.get("Jogo", ""),
            jogo.get("Nome", ""),
            jogo.get("Preço"),
            jogo.get("Prémio máximo", ""),
            jogo.get("Emissão", ""),
            jogo.get("Data lançamento", ""),
            jogo.get("contentId", ""),
            jogo.get("URL", ""),
            agora,
        ]

        ws.append(linha)

        celula_url = ws.cell(ws.max_row, 8)
        celula_url.hyperlink = celula_url.value
        celula_url.style = "Hyperlink"

        # Formata a coluna Preço como moeda (€)
        celula_preco = ws.cell(ws.max_row, 3)
        if celula_preco.value is not None:
            celula_preco.number_format = '#,##0.00 €'

    header_fill = PatternFill("solid", fgColor="808080")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    larguras = {
        "A": 10,
        "B": 32,
        "C": 12,
        "D": 18,
        "E": 18,
        "F": 22,
        "G": 14,
        "H": 75,
        "I": 22,
    }

    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura

    if ws.max_row >= 2:
        ref = f"A1:I{ws.max_row}"
        tabela = Table(displayName="TabelaRaspadinhas", ref=ref)
        estilo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tabela.tableStyleInfo = estilo
        ws.add_table(tabela)

    wb.save(caminho)


def main():
    print("=" * 60)
    print(" ATUALIZAÇÃO DAS RASPADINHAS - JOGOS SANTA CASA")
    print("=" * 60)

    with requests.Session() as session:
        session.headers.update(HEADERS)

        print("\n1. A consultar lista principal...")
        soup = pedir_pagina(session, URL_LISTA)

        jogos = extrair_jogos_da_lista(soup)

        if not jogos:
            raise RuntimeError(
                "Não foram encontradas raspadinhas. "
                "O site pode ter alterado a estrutura da página."
            )

        print(f"   Encontradas {len(jogos)} raspadinhas.")

        print("\n2. A consultar detalhes de cada raspadinha...")

        resultados = []

        for i, jogo in enumerate(jogos, start=1):
            print(f"   [{i:02d}/{len(jogos):02d}] contentId={jogo['contentId']}")

            try:
                resultado = extrair_detalhes(session, jogo)
                print(f"       -> Jogo {resultado.get('Jogo')} - {resultado.get('Nome')}")
                resultados.append(resultado)
            except Exception as erro:
                print(f"       ERRO: {erro}")
                resultados.append(jogo)

            time.sleep(0.3)

    print("\n3. A criar Excel...")
    criar_excel(resultados, FICHEIRO_EXCEL)

    print(f"\nConcluído!")
    print(f"Ficheiro criado: {FICHEIRO_EXCEL}")
    print(f"Total de raspadinhas: {len(resultados)}")


if __name__ == "__main__":
    main()
